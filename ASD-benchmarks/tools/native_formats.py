import sys
sys.dont_write_bytecode = True
"""Read-only structural inspection; never changes scientific payloads."""
import io,os,sys,json,pathlib,hashlib,zipfile,gzip,stat,csv,collections,tempfile,traceback,datetime,xml.etree.ElementTree as ET,wave
from audit_common import dump,digest,utc

def safe(name):
 q=pathlib.PurePosixPath(name.replace('\\','/'))
 return not(q.is_absolute() or '..' in q.parts or (q.parts and ':' in q.parts[0]))

def text_table(b,name):
 text=b.decode('utf-8-sig',errors='replace');lines=text.splitlines()
 out={'line_count':len(lines),'preview_lines':lines[:min(5,len(lines))]}
 if name.lower().endswith(('.csv','.tsv')):
  try:
   delim='\t' if name.lower().endswith('.tsv') else ','
   if text[:8192].count(';')>text[:8192].count(','):delim=';'
   rows=csv.reader(io.StringIO(text),delimiter=delim);head=next(rows,[]);counts=collections.Counter(len(r) for r in rows)
   out={'delimiter':delim,'header':head,'data_records':sum(counts.values()),'row_width_counts':dict(counts)}
  except csv.Error as e:out['csv_note']=str(e)
 return out

def inspect_bytes(b,name):
 ext=pathlib.PurePosixPath(name).suffix.lower();o={}
 if name.startswith('__MACOSX/') or pathlib.PurePosixPath(name).name.startswith('._'):return {'format':'source_OS_resource_fork'}
 if ext in ('.xlsx','.xlsm'):
  import openpyxl
  z=zipfile.ZipFile(io.BytesIO(b));assert z.testzip() is None
  w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=False)
  o={'format':'Excel_OpenXML','sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column,'first_rows':[[str(v)[:160] if v is not None else None for v in r] for r in list(__import__('itertools').islice(s.values,5))]} for s in w]};w.close()
 elif ext in ('.tif','.tiff','.jpg','.jpeg','.png','.gif'):
  from PIL import Image
  with Image.open(io.BytesIO(b)) as im:
   o={'format':im.format,'size':im.size,'mode':im.mode,'frames':getattr(im,'n_frames',1)};im.verify()
 elif ext=='.wav':
  with wave.open(io.BytesIO(b),'rb') as w:o={'format':'WAVE','channels':w.getnchannels(),'sample_rate_Hz':w.getframerate(),'sample_width_bytes':w.getsampwidth(),'frames':w.getnframes(),'duration_seconds':w.getnframes()/w.getframerate()};assert len(w.readframes(w.getnframes()))==w.getnframes()*w.getnchannels()*w.getsampwidth()
 elif ext=='.abf':
  import pyabf
  with tempfile.NamedTemporaryFile(suffix='.abf') as t:
   t.write(b);t.flush();a=pyabf.ABF(t.name,loadData=False)
   o={'format':'ABF','channels':a.channelCount,'sample_rate_Hz':a.dataRate,'sweeps':a.sweepCount,'units':a.adcUnits,'acquisition_datetime':str(a.abfDateTime)}
 elif ext=='.pdf':
  # pypdf is in bundled runtime, fallback to local PyPDF2.
  try:from pypdf import PdfReader
  except ImportError:from PyPDF2 import PdfReader
  r=PdfReader(io.BytesIO(b));o={'format':'PDF','pages':len(r.pages),'metadata':str(r.metadata)[:2000]}
 elif ext in ('.json','.ipynb'):
  j=json.loads(b);o={'format':'JSON','root_type':type(j).__name__,'top_level_keys':list(j)[:30] if isinstance(j,dict) else None,'length':len(j) if hasattr(j,'__len__') else None}
 elif ext=='.jsonl':
  keys=collections.Counter();n=0
  for line in b.splitlines():
   if line.strip():j=json.loads(line);keys.update(j.keys() if isinstance(j,dict) else []);n+=1
  o={'format':'JSONL','records':n,'key_counts':dict(keys)}
 elif ext=='.xml':
  root=ET.fromstring(b);o={'format':'XML','root':root.tag,'element_count':sum(1 for _ in root.iter()),'top_tags':[x.tag for x in list(root)[:15]]}
 elif ext in ('.csv','.tsv','.txt','.dat','.md','.hea','.m','.zpl','.py','.jws','.asc','.xmu','.cif','.xyn') or not ext:
  o={'format':'text_or_instrument_native','signature_hex':b[:12].hex()}
  if b'\x00' not in b[:2048]:o.update(text_table(b,name))
 else:o={'format':'preserved_native','signature_hex':b[:16].hex()}
 return o

def inspect_file(p):
 ext=p.suffix.lower();r={'bytes':p.stat().st_size,'sha256':digest(p),'format_extension':ext,'checks':[]}
 if ext=='.zip':
  with zipfile.ZipFile(p) as z:
   assert z.testzip() is None;r['checks'].append('all_member_CRCs')
   members=[];expanded=0;forms=collections.Counter()
   for f in z.infolist():
    assert safe(f.filename),f.filename
    assert not stat.S_ISLNK(f.external_attr>>16),f.filename
    expanded+=f.file_size
    if f.is_dir():continue
    e=pathlib.PurePosixPath(f.filename).suffix.lower();forms[e]+=1
    row={'name':f.filename,'bytes':f.file_size,'crc32':'%08x'%f.CRC}
    # Every well-supported member format is structurally inspected; preserve native exceptions.
    if e not in ('.zip','.7z'):
     try:row['structure']=inspect_bytes(z.read(f),f.filename)
     except Exception as ex:row['inspection_limitation']=type(ex).__name__+': '+str(ex)
    else:
     b=z.read(f)
     if e=='.zip':
      with zipfile.ZipFile(io.BytesIO(b)) as zz:
       assert zz.testzip() is None
       assert all(safe(x.filename) and not stat.S_ISLNK(x.external_attr>>16) for x in zz.infolist())
       row['nested_expanded_bytes']=sum(x.file_size for x in zz.infolist());expanded+=row['nested_expanded_bytes']-f.file_size
    members.append(row)
   r.update(format='ZIP',expanded_bytes=expanded,members=members,member_format_counts=dict(forms));r['checks'].append('safe_member_paths_no_symlinks')
 elif ext=='.7z':
  import py7zr
  with py7zr.SevenZipFile(p) as z:
   assert z.testzip() is None
  with py7zr.SevenZipFile(p) as z:
   fs=z.list();assert all(safe(f.filename) for f in fs)
   r.update(format='7Z',expanded_bytes=sum(f.uncompressed for f in fs),members=[{'name':f.filename,'bytes':f.uncompressed,'is_directory':f.is_directory} for f in fs],checks=['all_member_CRCs','safe_member_paths'])
 elif ext in ('.h5','.hdf5'):
  import h5py,numpy as np
  with h5py.File(p,'r') as f:
   ds=[]
   def visit(name,obj):
    if isinstance(obj,h5py.Dataset):
     # Force chunk decoding without modifying arrays.
     if obj.size:obj[tuple(slice(0,1) for _ in obj.shape)] if obj.shape else obj[()]
     ds.append({'name':name,'shape':obj.shape,'dtype':str(obj.dtype),'attributes':{k:str(v)[:500] for k,v in obj.attrs.items()}})
   f.visititems(visit);r.update(format='HDF5',datasets=ds,attributes={k:str(v)[:500] for k,v in f.attrs.items()})
 elif ext=='.npz':
  import numpy as np
  with zipfile.ZipFile(p) as z:assert z.testzip() is None
  ds=[]
  with np.load(p,allow_pickle=False) as f:
   for k in f.files:
    try:a=f[k];ds.append({'name':k,'shape':a.shape,'dtype':str(a.dtype)})
    except ValueError as ex:ds.append({'name':k,'note':str(ex)+'; pickle not executed'})
  r.update(format='NPZ',arrays=ds,checks=['all_member_CRCs','no_pickle_execution'])
 elif ext=='.nc':
  import netCDF4
  with netCDF4.Dataset(p) as f:
   f.set_auto_maskandscale(False);vs=[]
   for k,v in f.variables.items():
    if v.size:v[tuple(slice(0,1) for _ in v.shape)] if v.shape else v[...]
    vs.append({'name':k,'shape':v.shape,'dtype':str(v.dtype),'attributes':{a:str(v.getncattr(a))[:500] for a in v.ncattrs()}})
   r.update(format='NetCDF',dimensions={k:len(v) for k,v in f.dimensions.items()},variables=vs,attributes={k:str(f.getncattr(k))[:500] for k in f.ncattrs()})
 elif ext=='.cdf':
  import cdflib
  f=cdflib.CDF(str(p));info=f.cdf_info();vs=[]
  for k in info.zVariables+info.rVariables:
   vi=f.varinq(k);a=f.varattsget(k);row={'name':k,'last_record':vi.Last_Rec,'dimensions':vi.Dim_Sizes,'attributes':{k:str(v)[:500] for k,v in a.items()}}
   if vi.Last_Rec>=0:
    first=f.varget(k,startrec=0,endrec=0)
    if k=='Epoch':row['start_utc']=str(cdflib.cdfepoch.to_datetime(first));row['end_utc']=str(cdflib.cdfepoch.to_datetime(f.varget(k,startrec=vi.Last_Rec,endrec=vi.Last_Rec)))
   vs.append(row)
  r.update(format='NASA_CDF',variables=vs,attributes={k:str(v)[:500] for k,v in f.globalattsget().items()})
 elif ext=='.parquet':
  import pyarrow.parquet as pq
  f=pq.ParquetFile(p);rows=sum(batch.num_rows for batch in f.iter_batches());assert rows==f.metadata.num_rows
  r.update(format='Parquet',rows=rows,row_groups=f.num_row_groups,schema=str(f.schema_arrow),checks=['all_rows_decoded'])
 elif ext in ('.fit','.fits','.pha','.rsp','.rsp2'):
  from astropy.io import fits
  with fits.open(p,checksum=True) as f:
   f.verify('exception');hs=[]
   for h in f:
    hs.append({'name':h.name,'shape':getattr(h.data,'shape',None),'columns':h.columns.names if hasattr(h,'columns') else None,'units':list(h.columns.units) if hasattr(h,'columns') else None,'header':{k:str(h.header[k]) for k in ['DATE-OBS','DATE-END','TSTART','TSTOP','TELESCOP','INSTRUME','DETNAM','DATATYPE','BUNIT'] if k in h.header}})
   r.update(format='FITS',HDUs=hs)
 elif ext=='.gz':
  b=gzip.decompress(p.read_bytes());r.update(format='GZIP',expanded_bytes=len(b),inner=inspect_bytes(b,p.stem),checks=['gzip_CRC'])
 else:r.update(inspect_bytes(p.read_bytes(),p.name))
 return r
