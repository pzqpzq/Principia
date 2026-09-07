from __future__ import annotations

import base64
import csv
import gzip
import hashlib
import io
import zipfile
from pathlib import Path
from typing import Any

import numpy as np

from ..domain import DataAsset

MAX_PREVIEW_BYTES = 2 * 1024 * 1024
MAX_PREVIEW_ROWS = 20
MAX_PREVIEW_COLUMNS = 20


def _asset_path(asset: DataAsset, root: Path) -> Path:
    canonical = root.resolve(strict=True)
    relative = str(asset.metadata.get("relative_path") or "")
    path = (canonical / relative).resolve(strict=True)
    if not path.is_relative_to(canonical) or not path.is_file():
        raise ValueError("asset locator escapes its registered source")
    return path


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return str(value).replace("\x00", "")[:300]


def _rows(stream: io.TextIOBase, suffix: str) -> dict[str, Any]:
    sample = stream.read(32_000)
    stream.seek(0)
    delimiter = "\t" if suffix in {".tsv", ".tab"} else ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        pass
    records = []
    for row in csv.reader(stream, delimiter=delimiter):
        records.append([_cell(value) for value in row[:MAX_PREVIEW_COLUMNS]])
        if len(records) >= MAX_PREVIEW_ROWS:
            break
    return {"kind": "table", "rows": records, "truncated": True}


def _normalize_image(array: np.ndarray) -> bytes:
    from PIL import Image

    values = np.asarray(array)
    if values.ndim > 2:
        values = values[0]
    values = np.squeeze(values)
    if values.ndim != 2:
        side = max(1, int(np.sqrt(values.size)))
        values = values.ravel()[: side * side].reshape(side, side)
    finite = values[np.isfinite(values)]
    if not finite.size:
        scaled = np.zeros(values.shape, dtype=np.uint8)
    else:
        low, high = np.quantile(finite, [0.01, 0.99])
        if high <= low:
            high = low + 1
        scaled = np.clip((values - low) / (high - low), 0, 1)
        scaled = np.nan_to_num(scaled, nan=0.0)
        scaled = np.asarray(scaled * 255, dtype=np.uint8)
    image = Image.fromarray(scaled, mode="L")
    image.thumbnail((1024, 1024))
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _image_preview(path: Path, format_name: str) -> bytes | None:
    try:
        if format_name in {"png", "jpeg", "jpg"}:
            from PIL import Image

            with Image.open(path) as source:
                source.thumbnail((1024, 1024))
                clean = source.convert("RGB")
                output = io.BytesIO()
                clean.save(output, format="JPEG", quality=82, optimize=True)
                return output.getvalue()
        if format_name == "dicom":
            import pydicom

            dataset = pydicom.dcmread(path, force=True)
            return _normalize_image(np.asarray(dataset.pixel_array, dtype=float))
        if format_name == "fits":
            from astropy.io import fits

            with fits.open(path, memmap=True, lazy_load_hdus=True) as hdus:
                for hdu in hdus:
                    data = getattr(hdu, "data", None)
                    if data is not None and np.asarray(data).size:
                        try:
                            return _normalize_image(np.asarray(data, dtype=float))
                        except (TypeError, ValueError):
                            continue
        if format_name == "netcdf":
            import xarray as xr

            with xr.open_dataset(path, decode_cf=False, cache=False) as dataset:
                for value in dataset.data_vars.values():
                    if np.issubdtype(value.dtype, np.number) and value.size:
                        selection = {
                            dimension: slice(0, min(int(value.sizes[dimension]), 1024))
                            for dimension in value.dims
                        }
                        return _normalize_image(np.asarray(value.isel(selection).values, dtype=float))
        if format_name == "pptx":
            from PIL import Image, ImageOps

            with zipfile.ZipFile(path) as archive:
                members = sorted(
                    (
                        item
                        for item in archive.infolist()
                        if not item.is_dir()
                        and item.filename.startswith("ppt/media/")
                        and Path(item.filename).suffix.casefold()
                        in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
                        and item.file_size <= 32 * 1024 * 1024
                    ),
                    key=lambda item: item.filename,
                )[:4]
                thumbnails: list[Image.Image] = []
                for member in members:
                    with archive.open(member) as stream:
                        with Image.open(io.BytesIO(stream.read())) as source:
                            image = source.convert("RGB")
                            image.thumbnail((700, 500))
                            thumbnails.append(
                                ImageOps.expand(image.copy(), border=2, fill=(205, 211, 220))
                            )
                if thumbnails:
                    width = max(image.width for image in thumbnails)
                    height = sum(image.height for image in thumbnails) + 16 * (len(thumbnails) - 1)
                    canvas = Image.new("RGB", (width, height), "white")
                    offset = 0
                    for image in thumbnails:
                        canvas.paste(image, ((width - image.width) // 2, offset))
                        offset += image.height + 16
                    canvas.thumbnail((1024, 1024))
                    output = io.BytesIO()
                    canvas.save(output, format="JPEG", quality=84, optimize=True)
                    return output.getvalue()
    except Exception:
        return None
    return None


def build_preview(asset: DataAsset, root: Path) -> dict[str, Any]:
    """Return a bounded, path-free preview. DICOM metadata is never exposed."""

    path = _asset_path(asset, root)
    image_body = _image_preview(path, asset.format)
    if image_body is not None:
        if len(image_body) > MAX_PREVIEW_BYTES:
            raise ValueError("derived preview exceeds the 2 MiB response cap")
        mime = "image/jpeg" if asset.format in {"png", "jpeg", "jpg", "pptx"} else "image/png"
        digest = hashlib.sha256(image_body).hexdigest()
        return {
            "schema_version": "principia.data-preview/v1",
            "asset_id": asset.asset_id,
            "kind": "image",
            "mime_type": mime,
            "derived": asset.format not in {"png", "jpeg", "jpg"},
            "preview_sha256": digest,
            "data_url": f"data:{mime};base64,{base64.b64encode(image_body).decode()}",
            "warnings": list(asset.warnings),
        }
    table: dict[str, Any] | None = None
    try:
        if asset.format in {"csv", "tsv", "delimited", "delimited_text", "text"}:
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
                table = _rows(stream, path.suffix.casefold())
        elif asset.format == "delimited_gzip":
            with gzip.open(path, "rt", encoding="utf-8", errors="replace", newline="") as stream:
                table = _rows(stream, ".tsv")
        elif asset.format == "zip":
            with zipfile.ZipFile(path) as archive:
                members = [
                    item
                    for item in archive.infolist()
                    if not item.is_dir()
                    and Path(item.filename).suffix.casefold() in {".csv", ".tsv", ".txt"}
                    and item.file_size <= 64 * 1024 * 1024
                ]
                members.sort(key=lambda item: (item.filename, item.file_size))
                if members:
                    with archive.open(members[0]) as raw:
                        with io.TextIOWrapper(raw, encoding="utf-8-sig", errors="replace") as stream:
                            table = _rows(stream, Path(members[0].filename).suffix.casefold())
                    table["archive_member"] = members[0].filename
        elif asset.format == "xlsx":
            import openpyxl

            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                candidates: list[tuple[int, str, list[list[Any]]]] = []
                for sheet in workbook.worksheets:
                    rows: list[list[Any]] = []
                    occupied = 0
                    for row in sheet.iter_rows(
                        min_row=1,
                        max_row=min(int(sheet.max_row or 0), 80),
                        min_col=1,
                        max_col=min(int(sheet.max_column or 0), MAX_PREVIEW_COLUMNS),
                        values_only=True,
                    ):
                        values = [_cell(value) for value in row]
                        if any(value not in {None, ""} for value in values):
                            rows.append(values)
                            occupied += sum(value not in {None, ""} for value in values)
                        if len(rows) >= MAX_PREVIEW_ROWS:
                            break
                    if rows:
                        candidates.append((occupied, sheet.title, rows))
                if candidates:
                    _, title, rows = max(candidates, key=lambda item: (item[0], len(item[2])))
                    table = {
                        "kind": "table",
                        "rows": rows,
                        "truncated": True,
                        "worksheet": title,
                    }
            finally:
                workbook.close()
    except Exception:
        table = None
    if table is not None:
        table.update(
            {
                "schema_version": "principia.data-preview/v1",
                "asset_id": asset.asset_id,
                "derived": True,
                "warnings": list(asset.warnings),
            }
        )
        return table
    return {
        "schema_version": "principia.data-preview/v1",
        "asset_id": asset.asset_id,
        "kind": "metadata",
        "derived": False,
        "format": asset.format,
        "modality": asset.modality,
        "byte_size": asset.byte_size,
        "profile": dict(asset.metadata.get("profile") or {}),
        "warnings": [*asset.warnings, "No bounded content preview is available for this format."],
    }
