from __future__ import annotations

import csv
import gzip
import hashlib
import importlib.util
import json
import mimetypes
import os
import re
import stat
import tarfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import quote

from ..cancellation import check_cancelled
from ..domain import DataAsset, DataView, canonical_sha256
from .numeric_locale import detect_delimiter, looks_numeric

ADAPTER_VERSION = "1.4.2-4"
MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_ARCHIVE_BYTES = 1024 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 512 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 10_000
SAMPLE_BYTES = 256 * 1024
HASH_CHUNK_BYTES = 1024 * 1024
EXECUTABLE_SUFFIXES = {
    ".app",
    ".bat",
    ".bin",
    ".cmd",
    ".com",
    ".dll",
    ".dylib",
    ".exe",
    ".jar",
    ".msi",
    ".ps1",
    ".scr",
    ".sh",
    ".so",
}
NESTED_ARCHIVE_SUFFIXES = {".zip", ".tar", ".tgz", ".gz", ".bz2", ".xz", ".7z"}


@dataclass(frozen=True)
class Inspection:
    format: str
    modality: str
    mime_type: str
    adapter: str
    status: str
    metadata: dict[str, Any]
    warnings: list[str]


@dataclass(frozen=True)
class InventoryResult:
    assets: list[DataAsset]
    views: list[DataView]
    coverage: dict[str, Any]
    source_digest: str


def file_sha256_stream(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(HASH_CHUNK_BYTES):
            check_cancelled()
            digest.update(chunk)
    return digest.hexdigest()


def _optional(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def _safe_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _bounded_read(path: Path, limit: int = SAMPLE_BYTES) -> bytes:
    with path.open("rb") as handle:
        return handle.read(limit)


def _table_profile(text: str, *, suffix: str = "") -> dict[str, Any]:
    lines = [line for line in text.splitlines() if line.strip()][:101]
    if not lines:
        return {"sample_rows": 0, "columns": [], "column_count": 0}
    delimiter = detect_delimiter("\n".join(lines[:20]), suffix=suffix)
    rows = list(csv.reader(lines, delimiter=delimiter))
    width = max((len(row) for row in rows), default=0)
    first = rows[0][:500] if rows else []

    occupied = [value for value in first if value.strip()]
    numeric_fraction = (
        sum(looks_numeric(value) for value in occupied) / len(occupied) if occupied else 0.0
    )
    header_inferred = numeric_fraction == 0.0
    header = first if header_inferred else [f"column_{index + 1}" for index in range(width)]
    return {
        "delimiter": delimiter,
        "sample_rows": max(0, len(rows) - int(header_inferred)),
        "column_count": width,
        "columns": header,
        "header_inferred": header_inferred,
        "header_inference": (
            "source_header" if header_inferred else "synthetic_numeric_columns"
        ),
        "header_numeric_fraction": round(numeric_fraction, 3),
        "first_row_sample": first[:32],
        "sample_truncated": len(text.encode("utf-8", errors="ignore")) >= SAMPLE_BYTES,
    }


def _office_xml_text(raw: bytes, *, limit: int = 12_000) -> str:
    """Extract bounded visible Office text without evaluating relationships or macros."""

    text = raw.decode("utf-8", errors="replace")
    values = re.findall(r"<(?:a|w):t(?: [^>]*)?>(.*?)</(?:a|w):t>", text, flags=re.DOTALL)
    decoded = [
        value.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&#39;", "'")
        .replace("&quot;", '"')
        for value in values
    ]
    return " ".join(" ".join(decoded).split())[:limit]


def _matrix_market_profile(text: str) -> dict[str, Any]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    header = lines[0] if lines else ""
    dimensions: list[int] = []
    for line in lines[1:]:
        if line.startswith("%"):
            continue
        try:
            dimensions = [int(value) for value in line.split()[:3]]
        except ValueError:
            pass
        break
    return {"matrix_market_header": header, "dimensions": dimensions}


def _archive_member_path(name: str) -> PurePosixPath:
    normalized = name.replace("\\", "/")
    value = PurePosixPath(normalized)
    if value.is_absolute() or ".." in value.parts or not value.parts:
        raise ValueError("archive contains path traversal")
    return value


def _zip_profile(path: Path, *, allow_office_printer_settings: bool = False) -> dict[str, Any]:
    members: list[dict[str, Any]] = []
    total = 0
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_ARCHIVE_MEMBERS:
            raise ValueError("archive member count exceeds the 10,000-member limit")
        for info in infos:
            member = _archive_member_path(info.filename)
            mode = (info.external_attr >> 16) & 0xFFFF
            if stat.S_ISLNK(mode):
                raise ValueError("archive contains a symbolic link")
            if info.flag_bits & 0x1:
                raise ValueError("archive contains encrypted data")
            if info.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                raise ValueError("archive member exceeds the 512 MiB limit")
            total += info.file_size
            if total > MAX_ARCHIVE_BYTES:
                raise ValueError("archive logical size exceeds the 1 GiB limit")
            suffix = member.suffix.casefold()
            printer_setting = (
                allow_office_printer_settings
                and suffix == ".bin"
                and member.parts[:2] == ("xl", "printerSettings")
            )
            if suffix in EXECUTABLE_SUFFIXES and not printer_setting:
                raise ValueError("archive contains an unexpected executable")
            if suffix in NESTED_ARCHIVE_SUFFIXES:
                raise ValueError("nested archives are not supported")
            if len(members) < 1_000:
                members.append(
                    {
                        "path": member.as_posix(),
                        "bytes": info.file_size,
                        "compressed_bytes": info.compress_size,
                        "directory": info.is_dir(),
                    }
                )
    return {
        "member_count": len(infos),
        "logical_bytes": total,
        "members": members,
        "members_truncated": len(infos) > len(members),
    }


def _tar_profile(path: Path) -> dict[str, Any]:
    members: list[dict[str, Any]] = []
    total = 0
    count = 0
    with tarfile.open(path, "r:*") as archive:
        for info in archive:
            count += 1
            if count > MAX_ARCHIVE_MEMBERS:
                raise ValueError("archive member count exceeds the 10,000-member limit")
            member = _archive_member_path(info.name)
            if info.issym() or info.islnk():
                raise ValueError("archive contains a link")
            if not (info.isfile() or info.isdir()):
                raise ValueError("archive contains an unsupported special member")
            if info.size > MAX_ARCHIVE_MEMBER_BYTES:
                raise ValueError("archive member exceeds the 512 MiB limit")
            total += info.size
            if total > MAX_ARCHIVE_BYTES:
                raise ValueError("archive logical size exceeds the 1 GiB limit")
            suffix = member.suffix.casefold()
            if suffix in EXECUTABLE_SUFFIXES:
                raise ValueError("archive contains an unexpected executable")
            if suffix in NESTED_ARCHIVE_SUFFIXES:
                raise ValueError("nested archives are not supported")
            if len(members) < 1_000:
                members.append({"path": member.as_posix(), "bytes": info.size, "directory": info.isdir()})
    return {
        "member_count": count,
        "logical_bytes": total,
        "members": members,
        "members_truncated": count > len(members),
    }


class AdapterRegistry:
    """Read-only, metadata-first format registry used by the ASD data plane."""

    def inspect(self, path: Path) -> Inspection:
        size = path.stat().st_size
        lower = path.name.casefold()
        if lower == ".ds_store" or lower.startswith("._"):
            return Inspection(
                "filesystem_metadata",
                "context",
                "application/octet-stream",
                "filesystem-metadata",
                "metadata_only",
                {},
                ["operating-system metadata is retained but excluded from scientific analysis"],
            )
        if lower.startswith("~$"):
            return Inspection(
                "office_lock",
                "context",
                "application/octet-stream",
                "office-lock",
                "metadata_only",
                {},
                ["Office lock file is retained but excluded from scientific analysis"],
            )
        if size == 0:
            return Inspection(
                "empty",
                "context",
                "application/octet-stream",
                "metadata",
                "metadata_only",
                {},
                ["zero-byte file retained but not analyzable"],
            )
        if size > MAX_FILE_BYTES:
            return Inspection(
                "oversized",
                "context",
                "application/octet-stream",
                "streaming-metadata",
                "blocked",
                {"limit_bytes": MAX_FILE_BYTES},
                ["file exceeds the 200 MiB regular-file limit"],
            )

        raw = _bounded_read(path)
        suffix = path.suffix.casefold()
        mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        warnings: list[str] = []

        try:
            if raw.startswith(b"\x89HDF\r\n\x1a\n"):
                return self._hdf(path, lower)
            if raw.startswith(b"SIMPLE  =") or suffix in {".fits", ".fit", ".fts"}:
                return self._fits(path)
            if raw.startswith(b"root") or suffix == ".root":
                return self._root(path)
            if len(raw) > 132 and raw[128:132] == b"DICM" or suffix == ".dcm":
                return self._dicom(path)
            if raw.startswith(b"CDF") or suffix == ".nc":
                return self._netcdf(path)
            if raw.startswith(b"\x89PNG\r\n\x1a\n") or raw.startswith(b"\xff\xd8\xff"):
                return self._image(path, mime)
            if raw.startswith(b"PK\x03\x04"):
                return self._zip_container(path, suffix)
            if raw.startswith(b"\x1f\x8b"):
                return self._gzip(path, lower)
            if tarfile.is_tarfile(path):
                profile = _tar_profile(path)
                return Inspection("tar", "context", "application/x-tar", "safe-tar", "analyzable", profile, [])
            if raw.startswith(b"%PDF-"):
                return self._pdf(path)
            if raw.startswith(b"%%MatrixMarket"):
                return Inspection("matrix_market", "array", "text/plain", "matrix-market", "analyzable", _matrix_market_profile(_safe_text(raw)), [])
            if suffix == ".edf" or re.match(rb"[0-9 ]{8}.{72}.{80}", raw[:160], flags=re.DOTALL):
                return self._edf(path)
            if suffix in {".ras", ".rasx"}:
                return self._rigaku(path, suffix)
            if suffix in {".se", ".mod", ".scan", ".parms"}:
                return self._vendor_text(path, suffix, raw)
            if suffix == ".xlsx":
                return self._xlsx(path)
            if suffix == ".docx":
                return self._docx(path)
            if suffix == ".rtf" or raw.startswith(b"{\\rtf"):
                return Inspection("rtf", "context", "application/rtf", "rtf-text", "analyzable", {"sample_characters": len(_safe_text(raw))}, [])
            if suffix == ".pdf":
                return self._pdf(path)
            if suffix in {".json", ".jsonl", ".geojson", ".meta", ".schema"} or lower.endswith(".json"):
                return self._json(path, raw, mime)
            if suffix in {".xml", ".quakeml", ".html", ".htm"}:
                return Inspection(suffix.lstrip(".") or "xml", "context" if suffix in {".html", ".htm"} else "graph", mime, "xml-stream", "analyzable", {"sample_characters": len(_safe_text(raw))}, [])
            if suffix in {".csv", ".tsv", ".tab"} or b"\t" in raw[:4_096] or b"," in raw[:4_096]:
                profile = _table_profile(_safe_text(raw), suffix=suffix)
                # Content wins over an unfamiliar final suffix. Official bulk
                # tables such as BLS ``ce.data.*.Employment`` and
                # ``cu.data.*.AllItems`` are tabular even though the final
                # filename segment resembles an extension. Keeping their
                # format as ``delimited`` makes the projection executable
                # without renaming or rewriting the source file.
                detected_format = (
                    suffix.lstrip(".")
                    if suffix in {".csv", ".tsv", ".tab"}
                    else "delimited"
                )
                return Inspection(detected_format, "table", mime if mime != "application/octet-stream" else "text/plain", "delimited", "analyzable", profile, [])
            if suffix in {".txt", ".md", ".area", ".footnote", ".item", ".industry", ".series"} or not suffix:
                text = _safe_text(raw)
                if "\t" in text or "," in text:
                    profile = _table_profile(text, suffix=suffix)
                    return Inspection("delimited_text", "table", "text/plain", "delimited", "analyzable", profile, [])
                return Inspection("text", "context", "text/plain", "text", "analyzable", {"sample_characters": len(text)}, [])
        except (OSError, ValueError, zipfile.BadZipFile, tarfile.TarError) as exc:
            return Inspection(
                suffix.lstrip(".") or "unknown",
                "context",
                mime,
                "metadata",
                "blocked",
                {},
                [str(exc)[:500]],
            )

        warnings.append("no specialized adapter; retained as metadata only")
        return Inspection(suffix.lstrip(".") or "unknown", "context", mime, "metadata", "metadata_only", {}, warnings)

    def _gzip(self, path: Path, lower: str) -> Inspection:
        with gzip.open(path, "rb") as stream:
            raw = stream.read(SAMPLE_BYTES)
        text = _safe_text(raw)
        if ".mtx" in lower or raw.startswith(b"%%MatrixMarket"):
            return Inspection("matrix_market_gzip", "array", "application/gzip", "gzip-matrix-market", "analyzable", _matrix_market_profile(text), [])
        if ".json" in lower or text.lstrip().startswith(("{", "[")):
            return Inspection("json_gzip", "table", "application/gzip", "gzip-json", "analyzable", {"sample_characters": len(text), "sample_truncated": len(raw) == SAMPLE_BYTES}, [])
        profile = _table_profile(text, suffix=".tsv" if "\t" in text[:4_096] else ".csv")
        return Inspection("delimited_gzip", "table", "application/gzip", "gzip-delimited", "analyzable", profile, [])

    def _zip_container(self, path: Path, suffix: str) -> Inspection:
        if suffix == ".xlsx":
            return self._xlsx(path)
        if suffix == ".docx":
            return self._docx(path)
        if suffix == ".pptx":
            return self._pptx(path)
        if suffix == ".rasx":
            return self._rigaku(path, suffix)
        return Inspection("zip", "context", "application/zip", "safe-zip", "analyzable", _zip_profile(path), [])

    def _pptx(self, path: Path) -> Inspection:
        metadata = _zip_profile(path, allow_office_printer_settings=True)
        with zipfile.ZipFile(path) as archive:
            names = [item.filename for item in archive.infolist() if not item.is_dir()]
            slide_text: list[dict[str, Any]] = []
            for name in sorted(names):
                if not re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                    continue
                excerpt = _office_xml_text(archive.read(name), limit=2_000)
                if excerpt:
                    slide_text.append(
                        {
                            "slide": int(re.search(r"(\d+)", Path(name).stem).group(1)),
                            "text": excerpt,
                        }
                    )
                if len(slide_text) >= 40:
                    break
        metadata.update(
            {
                "slide_count": sum(
                    bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
                    for name in names
                ),
                "embedded_image_count": sum(
                    name.startswith("ppt/media/")
                    and Path(name).suffix.casefold()
                    in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
                    for name in names
                ),
                "slide_text": slide_text,
                "text_characters": sum(len(item["text"]) for item in slide_text),
            }
        )
        return Inspection(
            "pptx",
            "image",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "office-pptx",
            "analyzable",
            metadata,
            [],
        )

    def _docx(self, path: Path) -> Inspection:
        metadata = _zip_profile(path)
        with zipfile.ZipFile(path) as archive:
            text = (
                _office_xml_text(archive.read("word/document.xml"), limit=12_000)
                if "word/document.xml" in archive.namelist()
                else ""
            )
        metadata.update({"text_excerpt": text, "text_characters": len(text)})
        return Inspection(
            "docx",
            "context",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "office-docx",
            "analyzable",
            metadata,
            [],
        )

    def _json(self, path: Path, raw: bytes, mime: str) -> Inspection:
        profile: dict[str, Any] = {"sample_characters": len(_safe_text(raw))}
        if path.stat().st_size <= 8 * 1024 * 1024:
            try:
                value = json.loads(path.read_text(encoding="utf-8-sig"))
                profile.update(
                    {
                        "top_level": type(value).__name__,
                        "keys": list(value)[:500] if isinstance(value, dict) else [],
                        "item_count": len(value) if isinstance(value, (dict, list)) else 1,
                    }
                )
            except (OSError, UnicodeError, json.JSONDecodeError):
                profile["partial_json"] = True
        return Inspection("json", "table", mime if mime != "application/octet-stream" else "application/json", "json", "analyzable", profile, [])

    def _image(self, path: Path, mime: str) -> Inspection:
        profile: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("PIL"):
            from PIL import Image

            with Image.open(path) as image:
                profile = {"width": image.width, "height": image.height, "mode": image.mode}
        else:
            warnings.append("Pillow unavailable; image dimensions deferred")
        return Inspection(path.suffix.casefold().lstrip("."), "image", mime, "image", "analyzable", profile, warnings)

    def _fits(self, path: Path) -> Inspection:
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("astropy"):
            from astropy.io import fits

            with fits.open(path, memmap=True, lazy_load_hdus=True) as hdus:
                metadata["hdus"] = [
                    {
                        "index": index,
                        "name": str(hdu.name),
                        "shape": list(getattr(getattr(hdu, "data", None), "shape", ()) or ()),
                        "type": type(hdu).__name__,
                    }
                    for index, hdu in enumerate(hdus)
                ]
        else:
            warnings.append("Astropy unavailable; FITS HDUs deferred")
        return Inspection("fits", "array", "image/fits", "fits", "analyzable", metadata, warnings)

    def _hdf(self, path: Path, lower: str) -> Inspection:
        if lower.endswith(".nc"):
            return self._netcdf(path)
        format_name = "nwb" if lower.endswith(".nwb") else "hdf5"
        modality = "signal" if format_name == "nwb" else "array"
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("h5py"):
            import h5py

            datasets: list[dict[str, Any]] = []
            with h5py.File(path, "r", swmr=True) as handle:
                def visitor(name: str, value: Any) -> None:
                    if len(datasets) >= 500 or not isinstance(value, h5py.Dataset):
                        return
                    datasets.append(
                        {"path": name, "shape": list(value.shape), "dtype": str(value.dtype)}
                    )

                handle.visititems(visitor)
            metadata = {"datasets": datasets, "datasets_truncated": len(datasets) >= 500}
        else:
            warnings.append("h5py unavailable; dataset structure deferred")
        return Inspection(format_name, modality, "application/x-hdf5", format_name, "analyzable", metadata, warnings)

    def _root(self, path: Path) -> Inspection:
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("uproot"):
            import uproot

            with uproot.open(path) as handle:
                metadata["keys"] = [str(key) for key in handle.keys(recursive=True)[:500]]
        else:
            warnings.append("Uproot unavailable; ROOT trees deferred")
        return Inspection("root", "table", "application/x-root", "uproot", "analyzable", metadata, warnings)

    def _edf(self, path: Path) -> Inspection:
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("pyedflib"):
            import pyedflib

            reader = pyedflib.EdfReader(str(path))
            try:
                metadata = {
                    "channel_count": reader.signals_in_file,
                    "channels": reader.getSignalLabels()[:256],
                    "sample_frequencies": [
                        float(reader.getSampleFrequency(index))
                        for index in range(min(reader.signals_in_file, 256))
                    ],
                    "duration_seconds": float(reader.file_duration),
                }
            finally:
                reader.close()
        else:
            warnings.append("pyEDFlib unavailable; channel structure deferred")
        return Inspection("edf", "signal", "application/edf", "edf", "analyzable", metadata, warnings)

    def _dicom(self, path: Path) -> Inspection:
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("pydicom"):
            import pydicom

            dataset = pydicom.dcmread(path, stop_before_pixels=True, force=True)
            safe_tags = {
                "Modality",
                "Rows",
                "Columns",
                "NumberOfFrames",
                "SeriesInstanceUID",
                "SOPInstanceUID",
                "SliceThickness",
                "PixelSpacing",
                "ImagePositionPatient",
                "ImageOrientationPatient",
            }
            metadata = {
                key: str(getattr(dataset, key))[:500]
                for key in sorted(safe_tags)
                if hasattr(dataset, key)
            }
        else:
            warnings.append("pydicom unavailable; safe DICOM tags deferred")
        return Inspection("dicom", "volume", "application/dicom", "dicom", "analyzable", metadata, warnings)

    def _netcdf(self, path: Path) -> Inspection:
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("xarray"):
            import xarray as xr

            with xr.open_dataset(path, decode_cf=False, cache=False) as dataset:
                metadata = {
                    "dimensions": {str(key): int(value) for key, value in dataset.sizes.items()},
                    "variables": [str(value) for value in list(dataset.variables)[:500]],
                }
        else:
            warnings.append("xarray unavailable; NetCDF variables deferred")
        return Inspection("netcdf", "array", "application/x-netcdf", "netcdf", "analyzable", metadata, warnings)

    def _pdf(self, path: Path) -> Inspection:
        metadata: dict[str, Any] = {}
        warnings: list[str] = []
        if _optional("pypdf"):
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            metadata["page_count"] = len(reader.pages)
            excerpts: list[dict[str, Any]] = []
            for index, page in enumerate(reader.pages[:8], start=1):
                try:
                    text = " ".join((page.extract_text() or "").split())[:2_000]
                except Exception:
                    text = ""
                if text:
                    excerpts.append({"page": index, "text": text})
            metadata["page_text"] = excerpts
            metadata["text_characters"] = sum(len(item["text"]) for item in excerpts)
        else:
            warnings.append("pypdf unavailable; PDF pages deferred")
        return Inspection("pdf", "context", "application/pdf", "pdf", "analyzable", metadata, warnings)

    def _xlsx(self, path: Path) -> Inspection:
        metadata = _zip_profile(path, allow_office_printer_settings=True)
        warnings: list[str] = []
        if _optional("openpyxl"):
            import openpyxl

            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                sheets: list[dict[str, Any]] = []
                workbook_columns: list[str] = []
                for sheet in workbook.worksheets:
                    nonempty_rows = 0
                    nonempty_columns: set[int] = set()
                    sampled_rows: list[list[Any]] = []
                    for row_index, row in enumerate(
                        sheet.iter_rows(
                            min_row=1,
                            max_row=min(int(sheet.max_row or 0), 500),
                            min_col=1,
                            max_col=min(int(sheet.max_column or 0), 128),
                            values_only=True,
                        ),
                        start=1,
                    ):
                        occupied = [index for index, value in enumerate(row, start=1) if value is not None]
                        if occupied:
                            nonempty_rows = row_index
                            nonempty_columns.update(occupied)
                            if len(sampled_rows) < 12:
                                sampled_rows.append(
                                    [
                                        (str(value)[:160] if value is not None else "")
                                        for value in row[:64]
                                    ]
                                )
                    header_candidates: list[str] = []
                    if sampled_rows:
                        for candidate_row in sampled_rows[:4]:
                            occupied_values = [value for value in candidate_row if value]
                            numeric = 0
                            for value in occupied_values:
                                try:
                                    float(value.replace(",", ""))
                                    numeric += 1
                                except ValueError:
                                    pass
                            numeric_fraction = numeric / max(1, len(occupied_values))
                            if occupied_values and numeric_fraction < 0.5:
                                header_candidates = candidate_row
                                break
                    named_columns = [
                        f"{sheet.title}::{value}"
                        for value in header_candidates
                        if value
                    ][:128]
                    workbook_columns.extend(named_columns)
                    sheets.append(
                        {
                            "title": sheet.title,
                            "declared_max_row": int(sheet.max_row or 0),
                            "declared_max_column": int(sheet.max_column or 0),
                            "bounded_nonempty_row": nonempty_rows,
                            "bounded_nonempty_columns": len(nonempty_columns),
                            "header_candidates": header_candidates[:64],
                            "representative_cells": [
                                row[:32] for row in sampled_rows[:6]
                            ],
                            "profile_truncated": bool(
                                int(sheet.max_row or 0) > 500
                                or int(sheet.max_column or 0) > 128
                            ),
                        }
                    )
                metadata["sheets"] = sheets
                metadata["columns"] = list(dict.fromkeys(workbook_columns))[:1_000]
            finally:
                workbook.close()
        else:
            warnings.append("openpyxl unavailable; worksheet dimensions deferred")
        return Inspection("xlsx", "table", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "office-xlsx", "analyzable", metadata, warnings)

    def _rigaku(self, path: Path, suffix: str) -> Inspection:
        if suffix == ".rasx":
            profile = _zip_profile(path)
            return Inspection("rasx", "signal", "application/zip", "rigaku-rasx", "analyzable", profile, [])
        text = _safe_text(_bounded_read(path))
        header_count = sum(line.startswith("*") for line in text.splitlines())
        data_rows = 0
        for line in text.splitlines():
            values = line.strip().split()
            if len(values) >= 2:
                try:
                    float(values[0])
                    float(values[1])
                    data_rows += 1
                except ValueError:
                    pass
        return Inspection("ras", "signal", "text/plain", "rigaku-ras", "analyzable", {"sample_header_lines": header_count, "sample_numeric_rows": data_rows}, [])

    def _vendor_text(self, path: Path, suffix: str, raw: bytes) -> Inspection:
        text = _safe_text(raw)
        printable = sum(character.isprintable() or character.isspace() for character in text)
        ratio = printable / max(1, len(text))
        if ratio < 0.75:
            return Inspection(suffix.lstrip("."), "signal", "application/octet-stream", "vendor-metadata", "metadata_only", {"text_ratio": round(ratio, 3)}, ["opaque vendor representation; related exports remain analyzable"])
        profile = _table_profile(text, suffix=".tsv") if "\t" in text else {"sample_characters": len(text)}
        return Inspection(suffix.lstrip("."), "table", "text/plain", "vendor-text", "analyzable", profile, [])


class AssetInventory:
    def __init__(self, registry: AdapterRegistry | None = None) -> None:
        self.registry = registry or AdapterRegistry()

    @staticmethod
    def _provenance_roles(root: Path) -> tuple[dict[str, str], dict[str, dict[str, Any]]]:
        path = root / "PROVENANCE.json"
        if not path.is_file():
            return {}, {}
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {}, {}
        if payload.get("schema") != "principia.local-scenario/v1":
            return {}, {}
        roles: dict[str, str] = {}
        details: dict[str, dict[str, Any]] = {}
        for item in payload.get("files") or []:
            if not isinstance(item, dict) or not str(item.get("path") or ""):
                continue
            relative = str(item["path"]).replace("\\", "/")
            raw_role = str(item.get("role") or "")
            roles[relative] = "raw" if "payload" in raw_role else "context"
            details[relative] = item
        return roles, details

    def inventory(self, *, root: Path, source_id: str, study_id: str) -> InventoryResult:
        root = root.resolve(strict=True)
        if not root.is_dir():
            raise NotADirectoryError(root)
        provenance_roles, provenance_details = self._provenance_roles(root)
        paths: list[Path] = []
        for directory, names, filenames in os.walk(root, followlinks=False):
            check_cancelled()
            names.sort()
            filenames.sort()
            current = Path(directory)
            for name in names:
                candidate = current / name
                if candidate.is_symlink():
                    raise ValueError("source contains a symbolic-link directory")
            for name in filenames:
                candidate = current / name
                if candidate.is_symlink() or not candidate.is_file():
                    raise ValueError("source contains a non-regular file")
                paths.append(candidate)
        if len(paths) > MAX_ARCHIVE_MEMBERS:
            raise ValueError("source contains more than 10,000 files")

        assets: list[DataAsset] = []
        views: list[DataView] = []
        source_records: list[dict[str, Any]] = []
        for path in paths:
            check_cancelled()
            relative = path.relative_to(root).as_posix()
            digest = file_sha256_stream(path)
            inspection = self.registry.inspect(path)
            # USER_BRIEF is generated user context even if a provenance manifest
            # lists it as context; it can guide interpretation but is never evidence.
            role = "user_context" if relative == "USER_BRIEF.txt" else provenance_roles.get(relative)
            if not role:
                if relative in {"PROVENANCE.json", "SHA256SUMS"}:
                    role = "metadata"
                elif relative in {"SCENARIO.md", "CORPUS_README.md"}:
                    role = "context"
                elif relative.startswith("raw/"):
                    role = "raw"
                elif relative.startswith("context/"):
                    role = "context"
                elif inspection.format in {"filesystem_metadata", "office_lock"}:
                    role = "metadata"
                elif inspection.modality in {"table", "array", "signal", "volume", "graph"}:
                    # Real user folders rarely carry a provenance manifest or a
                    # raw/context hierarchy. Treat executable scientific payloads
                    # at the root as data, otherwise the discovery engine can
                    # truthfully inventory them yet silently execute zero tests.
                    role = "raw"
                elif inspection.modality == "image" and inspection.format != "pptx":
                    role = "raw"
                elif inspection.format == "xlsx":
                    role = "raw"
                elif inspection.format in {"pptx", "docx", "pdf", "rtf", "text"}:
                    role = "context"
                else:
                    role = "metadata"
            stable = canonical_sha256({"study": study_id, "source": source_id, "path": relative})[:24]
            asset = DataAsset(
                asset_id=f"asset:{stable}",
                study_id=study_id,
                source_id=source_id,
                portable_uri=f"{source_id}/{quote(relative, safe='/')}",
                byte_sha256=digest,
                byte_size=path.stat().st_size,
                format=inspection.format,
                role=role,
                modality=inspection.modality,  # type: ignore[arg-type]
                mime_type=inspection.mime_type,
                adapter=inspection.adapter,
                adapter_version=ADAPTER_VERSION,
                status=inspection.status,  # type: ignore[arg-type]
                warnings=inspection.warnings,
                metadata={
                    "relative_path": relative,
                    "profile": inspection.metadata,
                    "official_manifest": provenance_details.get(relative, {}),
                    "evidence_eligible": role not in {"user_context", "metadata"}
                    and inspection.format not in {"filesystem_metadata", "office_lock"},
                },
            )
            assets.append(asset)
            source_records.append(
                {
                    "path": relative,
                    "sha256": digest,
                    "bytes": asset.byte_size,
                    "adapter": asset.adapter,
                    "status": asset.status,
                }
            )
            if asset.status == "analyzable":
                content_digest = canonical_sha256(
                    {
                        "asset": digest,
                        "adapter": asset.adapter,
                        "adapter_version": asset.adapter_version,
                        "profile": inspection.metadata,
                    }
                )
                dimensions = {
                    str(key): int(value)
                    for key, value in (inspection.metadata.get("dimensions") or {}).items()
                } if isinstance(inspection.metadata.get("dimensions"), dict) else {}
                if not dimensions and isinstance(inspection.metadata.get("dimensions"), list):
                    dimensions = {
                        f"axis_{index}": int(value)
                        for index, value in enumerate(inspection.metadata["dimensions"][:8])
                    }
                variables = inspection.metadata.get("variables") or inspection.metadata.get("columns") or []
                views.append(
                    DataView(
                        view_id="view:"
                        + canonical_sha256(
                            {"study": study_id, "content_digest": content_digest}
                        )[:24],
                        study_id=study_id,
                        asset_ids=[asset.asset_id],
                        kind=asset.modality,
                        name=Path(relative).name,
                        dimensions=dimensions,
                        variables=[str(value)[:300] for value in list(variables)[:5_000]],
                        locator={"asset_id": asset.asset_id, "relative_path": relative},
                        profile=inspection.metadata,
                        content_digest=content_digest,
                    )
                )

        source_digest = canonical_sha256(source_records)
        counts: dict[str, int] = {}
        modalities: dict[str, int] = {}
        formats: dict[str, int] = {}
        for asset in assets:
            counts[asset.status] = counts.get(asset.status, 0) + 1
            modalities[asset.modality] = modalities.get(asset.modality, 0) + 1
            formats[asset.format] = formats.get(asset.format, 0) + 1
        coverage = {
            "asset_count": len(assets),
            "view_count": len(views),
            "status_counts": counts,
            "modality_counts": modalities,
            "format_counts": formats,
            "total_bytes": sum(asset.byte_size for asset in assets),
            "read_only": True,
        }
        return InventoryResult(assets=assets, views=views, coverage=coverage, source_digest=source_digest)
