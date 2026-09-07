from __future__ import annotations

import numpy as np
import pytest

from principia.data_discovery.expression_search import LawSearchSpec, search_expressions


def test_linear_law_is_tested_against_no_input_null():
    rng = np.random.default_rng(1024)
    groups = np.repeat(np.arange(30), 12)
    x = rng.uniform(0.1, 3, (len(groups), 1))
    y = 2 + 4 * x[:, 0] + rng.normal(0, 0.01, len(groups))
    spec = LawSearchSpec(target="y", inputs=["x"], independent_unit_kind="specimen", max_candidates=100)
    result = search_expressions(spec, x, y, groups, source_digest="a" * 64)
    linear = [c for c in result["finalists"] if c["baseline_kind"] == "constant"]
    assert linear and any(c["passed"] for c in linear)
    assert min(c["test_nrmse"] for c in linear) < 0.02


def test_unknown_llm_family_cannot_remove_common_library():
    from principia.data_discovery.expression_search import enumerate_bases
    from principia.data_discovery.law_ast import ast_digest
    spec = LawSearchSpec(target="y", inputs=["x"], independent_unit_kind="specimen")
    ordinary = {ast_digest(ast) for ast in enumerate_bases(spec)}
    hinted = {ast_digest(ast) for ast in enumerate_bases(spec.model_copy(update={"family_kinds": ["unrecognized_family"]}))}
    assert ordinary <= hinted
    assert len(ordinary) >= 18


def test_unbound_numeric_fit_remains_a_candidate(tmp_path):
    from principia.data_discovery.adapters import AssetInventory
    from principia.data_discovery.exploratory_search import explore_unbound_views
    from principia.data_discovery.rule_engine import outcomes_to_law_bundle
    from principia.domain import ScientificProgram, StudyBlueprint
    x = np.linspace(0.2, 3, 50)
    (tmp_path / "anonymous.csv").write_text("alpha,beta\n" + "\n".join(f"{v},{1+v*v}" for v in x))
    inventory = AssetInventory().inventory(root=tmp_path, source_id="source:test", study_id="study:test")
    blueprint = StudyBlueprint(study_id="study:test", blueprint_digest="a" * 64)
    outcomes, receipts = explore_unbound_views(study_id="study:test", blueprint=blueprint,
        assets_and_roots=[(a, tmp_path) for a in inventory.assets], completed_assets=set(), check_control=lambda: None)
    assert outcomes and sum(r["evaluated_count"] for r in receipts) > 100
    assert all(o.result.independent_unit_count == 0 and o.finding.status != "supported_candidate" for o in outcomes)
    program = ScientificProgram(program_id="program:test", study_id="study:test", track="predictive_closure", title="Proposed numeric expressions", scientific_intent="Test provisional fitting", program_digest="b" * 64)
    bundle = outcomes_to_law_bundle(study_id="study:test", program=program, outcomes=outcomes)
    assert bundle.laws and not any(law.gate_summary["passed"] for law in bundle.laws)


def test_delete_data_project_reclaims_owned_studies_and_preserves_sources(tmp_path):
    from principia.application import Principia
    from principia.domain import JobRecord
    product = Principia.open(working_directory=tmp_path / "working", cloud_root=tmp_path / "cloud")
    try:
        raw = tmp_path / "source"
        raw.mkdir()
        (raw / "keep.csv").write_text("x,y\n1,2\n")
        product.repository.register_source("src:keep", raw, "external://keep", "Keep", "External/API")
        session = product.research_sessions.create_data_discovery_home(source_ids=["src:keep"], title="Delete me", provider_profile_id="", model="")
        sid, study = session["session_id"], "study:delete-test"
        product.repository.save_job(JobRecord(job_id="job:delete-test", kind="data_discovery", state="succeeded", stage="synthesize", progress=1))
        product.repository.create_data_study({"study_id": study, "job_id": "job:delete-test", "session_id": sid, "source_ids": ["src:keep"], "state": "succeeded", "phase": "synthesize"})
        with product.repository.connect() as conn:
            conn.execute("INSERT INTO workspace_records VALUES (?,?,?,'law','rules','data',1,'{}','digest','now')", (sid, study, "law:delete"))
            alias = dict(conn.execute("SELECT * FROM research_sessions WHERE session_id=?", (sid,)).fetchone())
            alias["session_id"] = "session:legacy-alias"
            alias["canonical_session_id"] = sid
            conn.execute(f"INSERT INTO research_sessions ({','.join(alias)}) VALUES ({','.join('?' for _ in alias)})", tuple(alias.values()))
            conn.execute("INSERT INTO research_session_aliases VALUES (?,?,?,?,?,?)", (alias["session_id"], sid, "", alias["source_set_digest"], "duplicate_source_set", "now"))
        product.repository.save_job(JobRecord(job_id="job:alias-test", kind="data_discovery", state="succeeded", stage="synthesize", progress=1))
        product.repository.create_data_study({"study_id": "study:alias-test", "job_id": "job:alias-test", "session_id": alias["session_id"], "source_ids": ["src:keep"], "state": "succeeded", "phase": "synthesize"})
        artifacts = product.research_sessions.artifacts_root / "data-discovery" / study
        outputs = product.research_sessions.outputs_root / study
        for directory in (artifacts, outputs):
            directory.mkdir(parents=True)
            (directory / "generated.json").write_text("{}")
        # Never delete after a stale browser revision.
        with pytest.raises(ValueError, match="revision conflict"):
            product.research_sessions.delete_session(sid, expected_revision=99)
        assert artifacts.exists()
        result = product.research_sessions.delete_session(sid, expected_revision=1)
        assert result["deleted_study_count"] == 2
        assert result["artifact_cleanup_pending"] == 0
        assert not artifacts.exists() and not outputs.exists()
        assert (raw / "keep.csv").read_text() == "x,y\n1,2\n"
        with product.repository.connect() as conn:
            for table in ("data_studies", "data_discovery_runs", "workspace_records", "research_sessions", "v14_jobs"):
                assert conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] == 0
            assert not conn.execute("PRAGMA foreign_key_check").fetchall()
        assert product.repository.source_root("src:keep") == raw
        product.research_sessions.consolidate_data_sessions()
        with product.repository.connect() as conn:
            assert conn.execute("SELECT COUNT(*) FROM research_sessions").fetchone()[0] == 0
    finally:
        product.close()


def test_spatial_group_never_crosses_partitions_when_regimes_differ():
    from types import SimpleNamespace

    from principia.data_discovery.scientific_programs import _split_objects
    objects = [SimpleNamespace(object_id=f"map:{group}:{regime}", regime=regime,
        driver=np.arange(9) + group + 1, coordinates=np.column_stack([np.arange(9), np.zeros(9)]),
        locator={"independent_group": f"specimen:{group}"}) for group in range(15) for regime in ("a", "b")]
    manifest, partitions = _split_objects("study:test", objects)
    assert manifest.development_unit_count + manifest.validation_unit_count + manifest.test_unit_count == 15
    for group in range(15):
        assert sum(any(o.locator["independent_group"] == f"specimen:{group}" for o in items) for items in partitions.values()) == 1


def test_large_spatial_equation_replays_and_canonicalizes_across_groupings():
    from principia.data_discovery.law_ast import (
        associative_node,
        ast_digest,
        canonicalize,
        evaluate,
    )
    from principia.data_discovery.scientific_programs import _FittedModel, _model_equation
    from principia.domain import EquationNode

    rng = np.random.default_rng(8)
    count = 247
    model = _FittedModel("hierarchical_kernel", 0.3, tuple(f"feature:{i}" for i in range(count)),
                        rng.normal(size=count), rng.uniform(0.2, 2, count), rng.normal(0, 0.02, count), 0.5)
    design = rng.normal(size=(17, count))
    ast, parameters = _model_equation(model)
    expected = np.exp(np.clip(model.intercept + ((design - model.means) / model.scales) @ model.coefficients, -50, 50))
    np.testing.assert_allclose(evaluate(ast, {f"z_{i}": design[:, i] for i in range(count)}, parameters), expected, rtol=1e-12)
    assert canonicalize(canonicalize(ast)) == canonicalize(ast)
    leaves = [EquationNode(op="variable", symbol=f"x_{i}") for i in range(500)]
    forward = associative_node("add", leaves)
    differently_grouped = associative_node("add", [associative_node("add", leaves[i:i + 10]) for i in range(0, 500, 10)][::-1])
    assert ast_digest(forward) == ast_digest(differently_grouped)


def test_numeric_materialization_failure_is_isolated(tmp_path, monkeypatch):
    import csv

    from principia.data_discovery import exploratory_search
    from principia.data_discovery.adapters import AssetInventory
    from principia.domain import StudyBlueprint

    (tmp_path / "malformed.csv").write_text("x,y\n1,2\n")
    inventory = AssetInventory().inventory(root=tmp_path, source_id="source:test", study_id="study:test")
    def fail(*args):
        raise csv.Error("malformed record")
    monkeypatch.setattr(exploratory_search, "load_numeric", fail)
    outcomes, receipts = exploratory_search.explore_unbound_views(study_id="study:test",
        blueprint=StudyBlueprint(study_id="study:test", blueprint_digest="a" * 64),
        assets_and_roots=[(a, tmp_path) for a in inventory.assets], completed_assets=set(), check_control=lambda: None)
    assert not outcomes and receipts
    assert receipts[0]["evaluated_count"] == 0
    assert "Error" in receipts[0]["reason"]


def test_vector_identity_executes_all_gates_and_retains_counterexamples(tmp_path):
    import uproot

    from principia.data_discovery.adapters import AssetInventory
    from principia.data_discovery.collection_operators import _root_transverse_vector_outcomes
    from principia.data_discovery.rule_engine import outcomes_to_law_bundle
    from principia.domain import ScientificProgram

    rng = np.random.default_rng(21)
    for index in range(4):
        px, py = rng.normal(size=(2, 500))
        with uproot.recreate(tmp_path / f"period{index}.root") as handle:
            handle["events"] = {"met": np.hypot(px, py), "met_phi": np.arctan2(py, px), "met_mpx": px, "met_mpy": py}
    def execute():
        inventory = AssetInventory().inventory(root=tmp_path, source_id="source:test", study_id="study:test")
        outcomes = _root_transverse_vector_outcomes("study:test", inventory.assets, tmp_path)
        program = ScientificProgram(program_id="program:test", study_id="study:test", track="mechanism", title="Geometry", scientific_intent="Test vector geometry", program_digest="b" * 64)
        return outcomes, outcomes_to_law_bundle(study_id="study:test", program=program, outcomes=outcomes)
    outcomes, bundle = execute()
    assert len(bundle.laws) == 1 and bundle.laws[0].gate_summary["passed"]
    assert all(r.execution_state == "passed" for r in bundle.laws[0].gate_receipts)
    broken = tmp_path / "period3.root"
    with uproot.open(broken) as handle:
        values = handle["events"].arrays(library="np")
    with uproot.recreate(broken) as handle:
        handle["events"] = {key: values[key] * (1.5 if key == "met" else 1) for key in ("met", "met_phi", "met_mpx", "met_mpy")}
    outcomes, bundle = execute()
    assert len(bundle.laws) == 1 and not bundle.laws[0].gate_summary["passed"]
    assert outcomes[0].result.estimate["file_receipts"]
    assert any(r.execution_state == "failed" for r in bundle.laws[0].gate_receipts)


def test_spatial_footer_conditions_are_retained_without_numeric_summaries(tmp_path):
    import openpyxl

    from principia.data_discovery.adapters import AssetInventory
    from principia.data_discovery.scientific_programs import _extract_blocks

    path = tmp_path / "arbitrary.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["X", "Y", *[f"specimen-{i}" for i in range(6)]])
    for x in range(5):
        for y in range(5):
            sheet.append([x, y, *[100 + i + x*x + y for i in range(6)]])
    sheet.append([])
    sheet.append([None, None, "P1", "P1", "P1", "P2", "P2", "P2"])
    sheet.append([None, "Mean", 100, 101, 102, 103, 104, 105])
    workbook.save(path)
    asset = AssetInventory().inventory(root=tmp_path, source_id="s", study_id="s").assets[0]
    blocks = _extract_blocks(asset, path)
    assert len(blocks) == 1
    assert {item["label"] for item in blocks[0].field_strata.values()} == {"P1", "P2"}
    for row in range(2, 27):
        for column in range(3, 9):
            sheet.cell(row, column, float(sheet.cell(row, column).value) * 1.37)
    workbook.save(path)
    assert _extract_blocks(asset, path)[0].field_strata == blocks[0].field_strata


def test_smooth_response_features_use_only_inputs_and_preserve_row_order():
    from types import SimpleNamespace

    from principia.data_discovery.scientific_programs import _features

    rng = np.random.default_rng(52)
    coordinates = np.array([(x, y) for x in np.linspace(-1, 1, 7) for y in np.linspace(-1, 1, 7)])
    obj = SimpleNamespace(coordinates=coordinates, target=np.ones(49), driver=rng.uniform(1, 2, 49), regime="a")
    design, names = _features(obj, "orthogonal_transport_18", ("a",))
    permutation = rng.permutation(49)
    changed = SimpleNamespace(coordinates=coordinates[permutation], target=rng.uniform(0, 100, 49), driver=obj.driver[permutation], regime="a")
    other, other_names = _features(changed, "orthogonal_transport_18", ("a",))
    assert names == other_names
    np.testing.assert_allclose(design[permutation], other, atol=1e-12)


def test_exploration_requires_a_response_in_coordinate_only_tables(tmp_path):
    from principia.data_discovery.adapters import AssetInventory
    from principia.data_discovery.exploratory_search import explore_unbound_views
    from principia.domain import StudyBlueprint

    (tmp_path / "points.csv").write_text("No.,X,Y\n" + "\n".join(f"{i},{i/10},{i*i/100}" for i in range(30)))
    inventory = AssetInventory().inventory(root=tmp_path, source_id="source:test", study_id="study:test")
    blueprint = StudyBlueprint(study_id="study:test", blueprint_digest="a" * 64)
    args = dict(study_id="study:test", assets_and_roots=[(a, tmp_path) for a in inventory.assets], completed_assets=set(), check_control=lambda: None)
    outcomes, receipts = explore_unbound_views(blueprint=blueprint, **args)
    assert not outcomes and receipts
    outcomes, _ = explore_unbound_views(blueprint=blueprint.model_copy(update={"target_bindings": [{"variable": "Y"}]}), **args)
    assert outcomes
    assert outcomes[0].result.independent_unit_count == 0
    assert all("No." != binding["meaning"] for binding in outcomes[0].result.equation_variables)
