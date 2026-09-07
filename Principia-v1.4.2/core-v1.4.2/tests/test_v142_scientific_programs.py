from __future__ import annotations

import math
import time
from pathlib import Path

import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient

from principia.api import create_app
from principia.application import Principia
from principia.data_discovery.law_ast import (
    ast_digest,
    differentiate,
    evaluate,
    infer_dimension,
    numerical_derivative,
    render_latex,
)
from principia.data_discovery.scientific_programs import (
    compile_scientific_programs,
    execute_replicated_spatial_program,
)
from principia.domain import DataAsset, EquationNode, StudyBlueprint, canonical_sha256
from principia.domain.hashing import file_sha256


def test_equation_ast_is_canonical_executable_and_dimension_checked() -> None:
    x = EquationNode(op="variable", symbol="x")
    beta = EquationNode(op="parameter", symbol="beta")
    first = EquationNode(op="add", children=[x, beta])
    second = EquationNode(op="add", children=[beta, x])
    assert ast_digest(first) == ast_digest(second)
    assert np.allclose(evaluate(first, {"x": [1.0, 2.0]}, {"beta": 3.0}), [4.0, 5.0])
    assert "beta" in render_latex(first)
    assert infer_dimension(
        EquationNode(
            op="multiply",
            children=[x, EquationNode(op="variable", symbol="time")],
        ),
        {"x": {"length": 1}, "time": {"time": 1}},
    ) == {"length": 1.0, "time": 1.0}
    with pytest.raises(ValueError, match="dimensionally incompatible"):
        infer_dimension(first, {"x": {"length": 1}, "beta": {"time": 1}})


def test_analytic_and_numerical_equation_derivatives_agree() -> None:
    x = EquationNode(op="variable", symbol="x")
    equation = EquationNode(
        op="exp",
        children=[
            EquationNode(
                op="add",
                children=[
                    EquationNode(
                        op="multiply",
                        children=[EquationNode(op="parameter", symbol="a"), x],
                    ),
                    EquationNode(op="parameter", symbol="b"),
                ],
            )
        ],
    )
    values = {"x": np.asarray([-0.3, 0.2, 0.8])}
    parameters = {"a": 1.7, "b": -0.2}
    analytic = evaluate(differentiate(equation, "x"), values, parameters)
    numeric = numerical_derivative(equation, "x", values, parameters)
    assert np.allclose(analytic, numeric, rtol=1e-6, atol=1e-7)


def _write_spatial_workbook(
    path: Path,
    *,
    sheet_name: str = "Replicated fields",
    x_label: str = "X",
    y_label: str = "Y",
    row_order: list[int] | None = None,
    map_count: int = 16,
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    driver_x_column = map_count + 5
    driver_y_column = driver_x_column + 1
    driver_start_column = driver_x_column + 2
    sheet.cell(1, 1, "Thickness response maps")
    sheet.cell(1, driver_x_column, "Conductance driver maps")
    sheet.cell(4, 1, x_label)
    sheet.cell(4, 2, y_label)
    sheet.cell(4, driver_x_column, x_label)
    sheet.cell(4, driver_y_column, y_label)
    coordinates = [(x, y) for x in np.linspace(-1, 1, 5) for y in np.linspace(-1, 1, 5)]
    for map_index in range(map_count):
        label = f"unit-{map_index:02d}"
        sheet.cell(4, 3 + map_index, label)
        sheet.cell(4, driver_start_column + map_index, label)
    order = row_order or list(range(len(coordinates)))
    for row_index, coordinate_index in enumerate(order, start=5):
        x, y = coordinates[coordinate_index]
        sheet.cell(row_index, 1, float(x))
        sheet.cell(row_index, 2, float(y))
        sheet.cell(row_index, driver_x_column, float(x))
        sheet.cell(row_index, driver_y_column, float(y))
        radius_squared = x * x + y * y
        for map_index in range(map_count):
            driver = math.exp(
                0.2 * map_index
                + (0.35 + 0.04 * map_index) * x
                - 0.25 * y
                + 0.12 * math.sin((map_index + 1) * x * y)
            )
            normalized_driver = driver / math.exp(0.2 * map_index)
            target = math.exp(
                5.0
                + 0.11 * map_index
                - 0.20 * radius_squared
                + 0.08 * (x * x - y * y)
                + 0.75 * math.log(normalized_driver)
            )
            sheet.cell(row_index, 3 + map_index, target)
            sheet.cell(row_index, driver_start_column + map_index, driver)
    workbook.save(path)


def test_replicated_spatial_program_uses_locked_complete_units(tmp_path: Path) -> None:
    workbook_path = tmp_path / "renamed-uninformative.xlsx"
    _write_spatial_workbook(workbook_path)
    asset = DataAsset(
        asset_id="asset:spatial",
        study_id="study:spatial",
        source_id="source:spatial",
        portable_uri="local://asset/spatial",
        byte_sha256=file_sha256(workbook_path),
        byte_size=workbook_path.stat().st_size,
        format="xlsx",
        role="raw",
        modality="table",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        adapter="office",
        status="analyzable",
        metadata={"relative_path": workbook_path.name},
    )
    blueprint = StudyBlueprint(
        study_id="study:spatial",
        blueprint_digest=canonical_sha256({"study": "spatial"}),
        confidence=0.8,
        needs_user_confirmation=False,
    )
    predictive = compile_scientific_programs(
        study_id="study:spatial", blueprint=blueprint, principle_ids=[]
    )[1]
    bundle = execute_replicated_spatial_program(
        study_id="study:spatial",
        program=predictive,
        assets_and_roots=[(asset, tmp_path)],
    )
    assert len(bundle.laws) == 1
    assert len(bundle.split_manifests) == 1
    split = bundle.split_manifests[0]
    assert split.development_unit_count >= 3
    assert split.validation_unit_count >= 1
    assert split.test_unit_count >= 1
    law = bundle.laws[0]
    executed = {gate.metric: gate for gate in law.gate_receipts}
    assert executed["ast_replay"].execution_state == "passed"
    assert executed["ast_replay"].details["maximum_relative_error"] <= 1e-10
    assert executed["dimensional_plausibility"].details["inferred_output"] == {"native_target": 1.0}
    assert bundle.calibrations[0].fitted_parameters["equation_parameters"]["T_ref"] == 1.0
    gates = law.gate_summary["gates"]
    assert gates["heldout_target_mutation_invariant"] is True
    assert gates["test_baseline_dominance"] is True
    assert gates["negative_control"] is True
    assert bundle.calibrations[0].test_metrics["unit_count"] == split.test_unit_count
    assert bundle.outcomes[0].result.split_validation["split_manifest_id"] == split.split_manifest_id


def test_spatial_program_is_invariant_to_names_and_row_order(tmp_path: Path) -> None:
    original_path = tmp_path / "first.xlsx"
    transformed_path = tmp_path / "second-unrelated-name.xlsx"
    _write_spatial_workbook(original_path)
    _write_spatial_workbook(
        transformed_path,
        sheet_name="Renamed experimental sheet",
        x_label="X coordinate",
        y_label="Y coordinate",
        row_order=list(reversed(range(25))),
    )

    def execute(path: Path):
        asset = DataAsset(
            asset_id="asset:" + path.stem,
            study_id="study:metamorphic",
            source_id="source:metamorphic",
            portable_uri="local://asset/" + path.stem,
            byte_sha256=file_sha256(path),
            byte_size=path.stat().st_size,
            format="xlsx",
            role="raw",
            modality="table",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            adapter="office",
            status="analyzable",
            metadata={"relative_path": path.name},
        )
        blueprint = StudyBlueprint(
            study_id="study:metamorphic",
            blueprint_digest=canonical_sha256({"study": "metamorphic"}),
            confidence=0.8,
            needs_user_confirmation=False,
        )
        program = compile_scientific_programs(
            study_id="study:metamorphic", blueprint=blueprint, principle_ids=[]
        )[1]
        return execute_replicated_spatial_program(
            study_id="study:metamorphic",
            program=program,
            assets_and_roots=[(asset, tmp_path)],
        )

    original = execute(original_path)
    transformed = execute(transformed_path)
    assert original.split_manifests[0].assignments == transformed.split_manifests[0].assignments
    assert original.laws[0].canonical_ast_digest == transformed.laws[0].canonical_ast_digest
    assert original.calibrations[0].test_metrics == pytest.approx(
        transformed.calibrations[0].test_metrics, rel=1e-5, abs=1e-6
    )


def test_scientific_program_and_law_endpoints_restore_persisted_state(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source"
    source.mkdir()
    workbook = source / "fields.xlsx"
    _write_spatial_workbook(workbook)
    product = Principia.open(
        working_directory=tmp_path / "working",
        cloud_root=tmp_path / "cloud",
    )
    try:
        product.repository.register_source(
            "source:fields", source, "external://fields", "Fields", "External/API"
        )
        app = create_app(product, test_mode=True)
        client = TestClient(app)
        created = client.post(
            "/api/v1/data-discoveries",
            headers={"X-Principia-Session": app.state.session_token},
            json={
                "source_ids": ["source:fields"],
                "provider": "siliconflow",
                "egress_confirmed": False,
                "budget": "fast",
            },
        )
        assert created.status_code == 202
        study_id = created.json()["study_id"]
        for _ in range(250):
            state = client.get(f"/api/v1/data-discoveries/{study_id}").json()["state"]
            if state in {"succeeded", "partial", "failed", "cancelled"}:
                break
            time.sleep(0.04)
        assert state == "partial"

        programs = client.get(
            f"/api/v1/data-discoveries/{study_id}/programs"
        ).json()["items"]
        assert {item["track"] for item in programs} == {
            "mechanism",
            "predictive_closure",
        }
        laws = client.get(f"/api/v1/data-discoveries/{study_id}/laws").json()["items"]
        promoted = next(item for item in laws if item["evidence_tier"] == "internal_locked_validation")
        assert promoted["promoted"]
        assert promoted["canonical_ast_digest"] == ast_digest(
            EquationNode.model_validate(promoted["equation_ast"])
        )
        detail = client.get(
            f"/api/v1/data-discoveries/{study_id}/laws/{promoted['law_id']}"
        )
        assert detail.status_code == 200
        body = detail.json()
        assert body["calibrations"]
        assert body["evaluations"]
        assert body["candidate_frontier"]
        calibration = body["calibrations"][0]
        calibration_response = client.get(
            f"/api/v1/data-discoveries/{study_id}/laws/{promoted['law_id']}"
            f"/calibrations/{calibration['calibration_id']}"
        )
        assert calibration_response.status_code == 200
        assert calibration_response.json()["test_metrics"]["unit_count"] >= 1
        rule_page = client.get(f"/api/v1/data-discoveries/{study_id}/rules").json()
        assert rule_page["items"]
        projected = next(item for item in rule_page["items"] if item.get("law_id") == promoted["law_id"])
        assert projected["expression_latex"] == promoted["expression_latex"]
        assert projected["record_kind"] == "data_rule"
    finally:
        product.close()
